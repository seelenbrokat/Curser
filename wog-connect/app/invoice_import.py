from __future__ import annotations

import csv
import io
import json
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

_IDENT_RE = re.compile(r"\b(\d{18})\b")
_AMOUNT_RE = re.compile(r"-?\d+(?:[.,]\d+)?")

_IDENT_HEADERS = (
    "identcode",
    "ident code",
    "identificationcode",
    "identification code",
    "sendungsnummer",
    "sendungs-barcode",
    "sendungsbarcode",
    "parcel code",
    "parcelcode",
    "barcode",
    "sendungs barcode",
    "paketnummer",
)

_AMOUNT_HEADERS = (
    "chf nach rabattberechnung",
    "chf nach rabatt",
    "chf mit individuellen preisen",
    "betrag chf",
    "betrag",
    "amount chf",
    "amount",
    "preis chf",
    "preis",
    "total chf",
    "total",
    "verrechnungsbetrag",
)

_PRODUCT_HEADERS = (
    "produkt/dienstleistung",
    "produkt",
    "dienstleistung",
    "product",
    "service",
    "leistung",
    "leistungsbezeichnung",
    "artikel",
)

_INVOICE_HEADERS = (
    "rechnungsnummer",
    "rechnung",
    "invoice number",
    "invoice",
    "rechnungs-nr",
)

_PERIOD_HEADERS = (
    "rechnungsperiode",
    "verrechnungsperiode",
    "periode",
    "billing period",
    "aufgabedatum",
    "datum",
)

_REF_HEADERS = (
    "itemid",
    "item id",
    "sendingid",
    "sending id",
    "kundenreferenz",
    "referenz",
    "kundenreferenznummer",
    "customer reference",
    "sendungsreferenz",
    "auftragsnummer",
    "kundenreferenz sendung",
)


def normalize_identcode(raw: str | None) -> str:
    if not raw:
        return ""
    digits = re.sub(r"\D", "", str(raw).strip())
    return digits[:18] if len(digits) >= 18 else digits


def parse_amount_chf(raw: Any) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = str(raw).strip()
    if not text:
        return None
    text = text.replace("'", "").replace("\u2019", "").replace("CHF", "").replace("chf", "").strip()
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    m = _AMOUNT_RE.search(text)
    if not m:
        return None
    try:
        return float(m.group(0).replace(",", "."))
    except ValueError:
        return None


def amount_to_cents(amount: float | None) -> int | None:
    if amount is None:
        return None
    return int(round(amount * 100))


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("\n", " ").replace("\r", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def _pick_column(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for idx, header in enumerate(headers):
        h = _norm_header(header)
        if not h:
            continue
        for cand in candidates:
            if cand == h or cand in h:
                return idx
    return None


def _detect_columns(headers: list[str]) -> dict[str, int | None]:
    return {
        "identcode": _pick_column(headers, _IDENT_HEADERS),
        "amount": _pick_column(headers, _AMOUNT_HEADERS),
        "product": _pick_column(headers, _PRODUCT_HEADERS),
        "invoice": _pick_column(headers, _INVOICE_HEADERS),
        "period": _pick_column(headers, _PERIOD_HEADERS),
        "customer_ref": _pick_column(headers, _REF_HEADERS),
    }


def _cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _extract_identcode_from_row(row: list[Any], ident_idx: int | None) -> str:
    if ident_idx is not None:
        norm = normalize_identcode(_cell(row, ident_idx))
        if norm:
            return norm
    for cell in row:
        if cell is None:
            continue
        text = str(cell)
        m = _IDENT_RE.search(re.sub(r"\D", " ", text))
        if m:
            return m.group(1)
        norm = normalize_identcode(text)
        if len(norm) >= 18:
            return norm[:18]
    return ""


@dataclass
class ParsedBillingLine:
    identcode: str
    amount_cents: int
    product_label: str = ""
    invoice_number: str = ""
    billing_period: str = ""
    customer_ref: str = ""
    raw: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParseResult:
    lines: list[ParsedBillingLine]
    header_row: int
    columns: dict[str, int | None]
    invoice_number: str = ""
    billing_period: str = ""
    warnings: list[str] = field(default_factory=list)


def _iter_rows_xlsx(content: bytes) -> Iterator[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        yield list(row)
    wb.close()


def _iter_rows_csv(content: bytes) -> Iterator[list[Any]]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = content.decode("utf-8", errors="replace")
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for row in reader:
        yield row


def _score_header(headers: list[str]) -> tuple[int, dict[str, int | None]]:
    cols = _detect_columns([str(h) for h in headers])
    score = sum(1 for k in ("identcode", "amount") if cols.get(k) is not None)
    if cols.get("product") is not None:
        score += 1
    return score, cols


def parse_billing_file(content: bytes, filename: str) -> ParseResult:
    lower = (filename or "").lower()
    # Post Verarbeitungsnachweis (PvE) als PDF
    if lower.endswith(".pdf") or content[:4] == b"%PDF":
        from app.pve_pdf import parse_pve_pdf

        return parse_pve_pdf(content, filename=filename)
    if lower.endswith((".xlsx", ".xlsm", ".xltx")):
        rows = list(_iter_rows_xlsx(content))
    elif lower.endswith(".csv"):
        rows = list(_iter_rows_csv(content))
    else:
        try:
            rows = list(_iter_rows_xlsx(content))
        except Exception:
            rows = list(_iter_rows_csv(content))

    if not rows:
        raise ValueError("Datei ist leer")

    best_row = 0
    best_score = -1
    best_cols: dict[str, int | None] = {}
    scan_limit = min(len(rows), 20)
    for i in range(scan_limit):
        score, cols = _score_header(rows[i])
        if score > best_score:
            best_score = score
            best_row = i
            best_cols = cols

    if best_score < 2:
        raise ValueError(
            "Keine passenden Spalten gefunden (Identcode + Betrag). "
            "Bitte Excel/CSV aus dem Post-Rechnungsmanager verwenden."
        )

    headers = [str(h or "") for h in rows[best_row]]
    data_rows = rows[best_row + 1 :]
    warnings: list[str] = []
    if best_cols.get("amount") is None:
        warnings.append("Betragsspalte nicht erkannt – es wird nach CHF-Werten in der Zeile gesucht.")

    default_invoice = ""
    default_period = ""
    inv_idx = best_cols.get("invoice")
    per_idx = best_cols.get("period")
    for row in data_rows[:30]:
        if inv_idx is not None and not default_invoice:
            default_invoice = str(_cell(row, inv_idx) or "").strip()
        if per_idx is not None and not default_period:
            default_period = str(_cell(row, per_idx) or "").strip()

    parsed: list[ParsedBillingLine] = []
    ident_idx = best_cols.get("identcode")
    amount_idx = best_cols.get("amount")
    product_idx = best_cols.get("product")
    ref_idx = best_cols.get("customer_ref")

    for row in data_rows:
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        ident = _extract_identcode_from_row(row, ident_idx)
        customer_ref = str(_cell(row, ref_idx) or "").strip()
        amount = parse_amount_chf(_cell(row, amount_idx))
        if amount is None:
            for cell in row:
                if cell is None:
                    continue
                text = str(cell)
                if "chf" in text.lower() or re.search(r"\d+[.,]\d{2}", text):
                    amount = parse_amount_chf(cell)
                    if amount is not None:
                        break
        if amount is None or amount == 0:
            continue
        if not ident and not customer_ref:
            continue
        cents = amount_to_cents(amount)
        if cents is None:
            continue
        product = str(_cell(row, product_idx) or "").strip()
        invoice_no = str(_cell(row, inv_idx) or default_invoice or "").strip()
        period = str(_cell(row, per_idx) or default_period or "").strip()
        raw = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}
        parsed.append(
            ParsedBillingLine(
                identcode=ident,
                amount_cents=cents,
                product_label=product,
                invoice_number=invoice_no,
                billing_period=period,
                customer_ref=customer_ref,
                raw=raw,
            )
        )

    if not parsed:
        raise ValueError("Keine Rechnungszeilen mit Identcode und Betrag gefunden.")

    return ParseResult(
        lines=parsed,
        header_row=best_row + 1,
        columns=best_cols,
        invoice_number=default_invoice,
        billing_period=default_period,
        warnings=warnings,
    )


def is_surcharge_hint(product_label: str) -> bool:
    text = (product_label or "").lower()
    hints = (
        "sperrgut",
        "bulky",
        "übergewicht",
        "uebergewicht",
        "nachverrechn",
        "zuschlag",
        "mehrkosten",
        "retour",
        "unzustell",
        "manuell",
        " + man",
        " + zfz",
        " + le1",
        " + le2",
        " + xxl",
        " + si",
        " + fra",
        "zustellzeit",
    )
    return any(h in text for h in hints)


def format_chf(cents: int | None) -> str:
    if cents is None:
        return "—"
    return f"{cents / 100:.2f}"
